# -*- coding: utf-8 -*-
"""Contact management view using PySide6."""

from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QGroupBox,
    QTableWidget, QTableWidgetItem, QLineEdit, QPushButton,
    QLabel, QHeaderView, QSplitter, QFormLayout, QFileDialog
)
from PySide6.QtCore import Qt
from typing import List, Optional
import csv
import os

from views.base import (
    create_table, update_table, create_button, create_input,
    create_label, create_separator, show_error, show_toast, confirm,
    DEFAULT_FONT
)
from models.contact import Contact, ContactRepository


class ContactView(QWidget):
    """Contact management widget with CRUD operations."""
    
    def __init__(self, parent: Optional[QWidget] = None):
        super().__init__(parent)
        self._contact_id: Optional[int] = None  # Current editing contact ID
        self._setup_ui()
        self._connect_signals()
        self._load_data()
    
    def _setup_ui(self) -> None:
        """Setup the UI layout."""
        main_layout = QHBoxLayout(self)
        
        # Left side: Form
        form_group = QGroupBox('联系人信息')
        form_layout = QVBoxLayout(form_group)
        
        # Form fields
        fields_layout = QFormLayout()
        self._name_input = create_input(placeholder='输入姓名')
        self._email_input = create_input(placeholder='输入邮箱地址')
        fields_layout.addRow('姓名:', self._name_input)
        fields_layout.addRow('邮箱:', self._email_input)
        form_layout.addLayout(fields_layout)
        
        form_layout.addStretch()
        form_layout.addWidget(create_separator())
        
        # Form buttons
        btn_layout = QHBoxLayout()
        self._save_btn = create_button('保存', primary=True)
        self._clear_btn = create_button('清空')
        btn_layout.addWidget(self._save_btn)
        btn_layout.addWidget(self._clear_btn)
        btn_layout.addStretch()
        form_layout.addLayout(btn_layout)
        
        # Right side: Table
        table_group = QGroupBox('联系人列表')
        table_layout = QVBoxLayout(table_group)
        
        self._table = create_table(
            headings=['ID', '姓名', '邮箱'],
            data=[]
        )
        self._table.setColumnWidth(0, 50)
        self._table.setColumnWidth(1, 120)
        table_layout.addWidget(self._table)
        
        # Table buttons
        table_btn_layout = QHBoxLayout()
        self._import_btn = create_button('导入')
        self._edit_btn = create_button('编辑')
        self._delete_btn = create_button('删除', danger=True)
        table_btn_layout.addWidget(self._import_btn)
        table_btn_layout.addWidget(self._edit_btn)
        table_btn_layout.addWidget(self._delete_btn)
        table_btn_layout.addStretch()
        table_layout.addLayout(table_btn_layout)
        
        # Add to main layout with splitter
        splitter = QSplitter(Qt.Horizontal)
        splitter.addWidget(form_group)
        splitter.addWidget(table_group)
        splitter.setSizes([250, 550])
        main_layout.addWidget(splitter)
    
    def _connect_signals(self) -> None:
        """Connect button signals to handlers."""
        self._save_btn.clicked.connect(self._on_save)
        self._clear_btn.clicked.connect(self._on_clear)
        self._import_btn.clicked.connect(self._on_import)
        self._edit_btn.clicked.connect(self._on_edit)
        self._delete_btn.clicked.connect(self._on_delete)
        self._table.itemSelectionChanged.connect(self._on_selection_changed)
    
    def _load_data(self) -> None:
        """Load contacts from database and update table."""
        contacts = ContactRepository.get_all()
        data = [[c.id, c.name, c.email] for c in contacts]
        update_table(self._table, data)
    
    def refresh(self) -> None:
        """Public method to refresh the view data."""
        self._load_data()
    
    def _on_save(self) -> None:
        """Handle save button click."""
        name = self._name_input.text().strip()
        email = self._email_input.text().strip()
        
        if not name:
            show_error('请输入姓名', parent=self)
            return
        if not email:
            show_error('请输入邮箱', parent=self)
            return
        
        try:
            if self._contact_id:
                ContactRepository.update(self._contact_id, name, email)
                show_toast('联系人已更新', parent=self)
            else:
                ContactRepository.create(name, email)
                show_toast('联系人已添加', parent=self)
            
            self._load_data()
            self._on_clear()
        except ValueError as e:
            show_error(str(e), parent=self)
        except Exception as e:
            show_error(f'保存失败: {e}', parent=self)
    
    def _on_clear(self) -> None:
        """Clear form fields."""
        self._contact_id = None
        self._name_input.clear()
        self._email_input.clear()
    
    def _on_edit(self) -> None:
        """Handle edit button click."""
        row = self._table.currentRow()
        if row < 0:
            show_error('请先选择要编辑的联系人', parent=self)
            return
        
        self._contact_id = int(self._table.item(row, 0).text())
        self._name_input.setText(self._table.item(row, 1).text())
        self._email_input.setText(self._table.item(row, 2).text())
    
    def _on_delete(self) -> None:
        """Handle delete button click."""
        row = self._table.currentRow()
        if row < 0:
            show_error('请先选择要删除的联系人', parent=self)
            return
        
        contact_id = int(self._table.item(row, 0).text())
        name = self._table.item(row, 1).text()
        
        if not confirm(f'确定要删除联系人 "{name}" 吗？', parent=self):
            return
        
        try:
            ContactRepository.delete(contact_id)
            show_toast('联系人已删除', parent=self)
            self._load_data()
            self._on_clear()
        except Exception as e:
            show_error(f'删除失败: {e}', parent=self)
    
    def _on_selection_changed(self) -> None:
        """Handle table selection change."""
        pass  # Can be used for preview functionality
    
    def _on_import(self) -> None:
        """Handle import button click - import contacts from CSV/Excel."""
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            '选择导入文件',
            '',
            'CSV 文件 (*.csv);;Excel 文件 (*.xlsx);;所有文件 (*.*)'
        )
        
        if not file_path:
            return
        
        try:
            imported_count = 0
            skipped_count = 0
            
            if file_path.lower().endswith('.csv'):
                imported_count, skipped_count = self._import_from_csv(file_path)
            elif file_path.lower().endswith('.xlsx'):
                imported_count, skipped_count = self._import_from_xlsx(file_path)
            else:
                show_error('不支持的文件格式，请使用 CSV 或 XLSX 文件', parent=self)
                return
            
            self._load_data()
            show_toast(f'导入完成: 成功 {imported_count} 条, 跳过 {skipped_count} 条', parent=self)
        except Exception as e:
            show_error(f'导入失败: {e}', parent=self)
    
    def _import_from_csv(self, file_path: str) -> tuple:
        """Import contacts from CSV file.
        
        Expected format: name,email (with or without header)
        Returns: (imported_count, skipped_count)
        """
        imported = 0
        skipped = 0
        
        with open(file_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            for row in reader:
                if len(row) < 2:
                    skipped += 1
                    continue
                
                name = row[0].strip()
                email = row[1].strip()
                
                # Skip header row
                if name.lower() in ('name', '姓名', 'names') and email.lower() in ('email', '邮箱', 'emails', 'mail'):
                    continue
                
                if not name or not email:
                    skipped += 1
                    continue
                
                try:
                    # Check if contact already exists
                    existing = ContactRepository.get_by_name(name)
                    if existing:
                        skipped += 1
                        continue
                    
                    ContactRepository.create(name, email)
                    imported += 1
                except Exception:
                    skipped += 1
        
        return imported, skipped
    
    def _import_from_xlsx(self, file_path: str) -> tuple:
        """Import contacts from Excel file.
        
        Expected format: First column = name, Second column = email
        Returns: (imported_count, skipped_count)
        """
        try:
            from openpyxl import load_workbook
        except ImportError:
            show_error('导入 Excel 需要安装 openpyxl: pip install openpyxl', parent=self)
            return 0, 0
        
        imported = 0
        skipped = 0
        
        wb = load_workbook(file_path, read_only=True)
        ws = wb.active
        
        first_row = True
        for row in ws.iter_rows(values_only=True):
            if len(row) < 2:
                skipped += 1
                continue
            
            name = str(row[0]).strip() if row[0] else ''
            email = str(row[1]).strip() if row[1] else ''
            
            # Skip header row
            if first_row:
                first_row = False
                if name.lower() in ('name', '姓名', 'names') and email.lower() in ('email', '邮箱', 'emails', 'mail'):
                    continue
            
            if not name or not email:
                skipped += 1
                continue
            
            try:
                existing = ContactRepository.get_by_name(name)
                if existing:
                    skipped += 1
                    continue
                
                ContactRepository.create(name, email)
                imported += 1
            except Exception:
                skipped += 1
        
        wb.close()
        return imported, skipped
